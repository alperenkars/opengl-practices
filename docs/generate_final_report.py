from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from pathlib import Path


DOCS_DIR = Path(__file__).resolve().parent
OUTPUT_PATH = DOCS_DIR / "final_report.docx"

doc = Document()

for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)


def add_paragraph(text="", bold=False, italic=False, size=12,
                  align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_heading(text, level=1):
    sizes = {1: 14, 2: 13, 3: 12}
    befores = {1: 14, 2: 10, 3: 8}
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(befores.get(level, 10))
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(sizes.get(level, 12))
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_separator():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Cm(0.75)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_code(code_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Cm(0.4)

    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F2F2F2")
    pPr.append(shd)

    pBdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "4")
        b.set(qn("w:color"), "BFBFBF")
        pBdr.append(b)
    pPr.append(pBdr)

    run = p.add_run(code_text)
    run.font.name = "Consolas"
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Consolas")
    rFonts.set(qn("w:hAnsi"), "Consolas")
    rFonts.set(qn("w:cs"), "Consolas")
    rPr.append(rFonts)
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    r.font.italic = True
    r.font.color.rgb = RGBColor(80, 80, 80)


def add_screenshot_placeholder(label):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell = table.cell(0, 0)
    cell.width = Cm(14)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for _ in range(7):
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
    centre = cell.paragraphs[3]
    centre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = centre.add_run("[ Screenshot unavailable in this environment ]")
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.italic = True
    r.font.color.rgb = RGBColor(120, 120, 120)
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "FAFAFA")
    tcPr.append(shd)
    add_caption(label)


def style_table(table, headers):
    for i, cell in enumerate(table.rows[0].cells):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.clear()
        run = p.add_run(headers[i])
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "D9D9D9")
        tcPr.append(shd)


def fill_table_row(row, values):
    for i, cell in enumerate(row.cells):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.clear()
        run = p.add_run(values[i])
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0, 0, 0)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)


title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.space_after = Pt(4)
r = title.add_run("Final Report")
r.font.name = "Times New Roman"
r.font.size = Pt(16)
r.font.bold = True
r.font.color.rgb = RGBColor(0, 0, 0)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(2)
r = subtitle.add_run("COMP410 Computer Graphics — Spring 2026")
r.font.name = "Times New Roman"
r.font.size = Pt(12)
r.font.italic = True
r.font.color.rgb = RGBColor(0, 0, 0)

for line in [
    ("Project Title:", "Virtual Koç University Campus Tour"),
    ("Team Members:", "Alperen Kars, Uğur Öner"),
    ("Instructor:", "Yücel Yemez"),
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r1 = p.add_run(line[0] + " ")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(12)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0, 0, 0)
    r2 = p.add_run(line[1])
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(12)
    r2.font.color.rgb = RGBColor(0, 0, 0)

add_separator()

add_heading("1. Introduction")
add_paragraph(
    "This report describes the final version of our COMP410 term project, the Virtual Koç "
    "University Campus Tour. The goal of the project was simple from the start. We wanted a "
    "user to open one program, move inside a recognisable 3D campus, and explore it from a "
    "first-person point of view with only C++ and OpenGL."
)
add_paragraph(
    "The final build now includes the full campus scene, textured and procedural materials, "
    "first-person and fly navigation, terrain following, building collision, a procedural "
    "skybox, day and night lighting, landmark shortcuts, and a small HUD with a minimap. "
    "The rest of this report explains how we built that pipeline from data collection to the "
    "final interactive demo."
)
add_screenshot_placeholder("Figure 1 — Final daytime view of the campus from an overview camera.")

add_heading("2. Tech Stack")
add_paragraph(
    "We kept the toolchain small and close to the course setup. The main project code is in "
    "C++17, and all rendering goes through the OpenGL 4.1 core pipeline."
)
headers = ["Component", "Version / Notes", "Role in the project"]
rows = [
    ["Language", "C++17", "Core application and render loop."],
    ["Graphics API", "OpenGL 4.1 core", "All rendering and GPU state management."],
    ["Shading language", "GLSL 410", "Vertex, fragment, skybox, and HUD shaders."],
    ["Math", "GLM", "Vectors, matrices, camera math, and transforms."],
    ["Windowing & input", "GLFW", "Window creation, keyboard, mouse, and callbacks."],
    ["Extension loader", "GLEW", "Loads the OpenGL functions we use."],
    ["Model loading", "Assimp", "Loads campus, Odeon, and clock tower meshes at runtime."],
    ["Image loading", "stb_image", "Loads diffuse textures and embedded images."],
    ["Build system", "GNU Make", "Builds the final `campus` binary."],
    ["External tools", "Blender + blosm + OSM", "Used for campus data preparation before export."],
]
t = doc.add_table(rows=1 + len(rows), cols=3)
t.style = "Table Grid"
style_table(t, headers)
for i, row in enumerate(rows):
    fill_table_row(t.rows[i + 1], row)

add_paragraph(
    "The build and run steps stayed simple during the whole project:",
    space_before=8, space_after=4
)
add_code(
    "brew install glm glfw glew assimp\n"
    "make\n"
    "./campus"
)

add_heading("3. Project Structure")
add_paragraph(
    "The repository is still easy to follow. The runtime code is in `src/`, shaders are under "
    "`shaders/`, and all models and textures live under `assets/`."
)
add_code(
    "src/          main loop, camera, mesh loading, texture loading\n"
    "shaders/      main shading, skybox pass, overlay pass\n"
    "assets/       campus model, extra models, facade textures, material textures\n"
    "docs/         proposal, progress report, final report scripts"
)

add_heading("4. Final Feature Summary")
add_paragraph(
    "By the final stage, the application covers both the core graphics topics and the extra "
    "demo features we wanted."
)
for item in [
    "Runtime loading of the campus scene from `Campus.glb` with Assimp.",
    "Loading of extra landmark models such as Odeon and the clock tower.",
    "Per-fragment lighting with ambient, diffuse, and Blinn-Phong specular terms.",
    "Support for both textured meshes and procedural material shading.",
    "Gamma-correct lighting and texture handling.",
    "First-person walking mode with mouse look and WASD movement.",
    "Fly mode for fast aerial inspection of the campus.",
    "Terrain following, safe spawn placement, and collision against buildings.",
    "Procedural skybox with a day and night lighting toggle.",
    "On-screen HUD with minimap, landmark labels, and status text.",
    "Landmark jump keys for Rectorate, Library, and Student Center.",
    "Per-mesh frustum culling to skip geometry outside the camera view.",
]:
    add_bullet(item)

add_heading("5. Asset Creation and Campus Pipeline")
add_paragraph(
    "The project started with data collection. We first took campus layout data from "
    "OpenStreetMap. That gave us the building footprints and the rough campus structure."
)
add_paragraph(
    "We then moved to Blender. Using blosm and manual cleanup, we turned the OSM data into a "
    "usable scene. We fixed mesh names, adjusted heights, prepared UVs where needed, and "
    "exported the result as a runtime-friendly campus file."
)
for item in [
    "OpenStreetMap gave us the base layout and building positions.",
    "Blender was used to clean the scene and prepare the final export.",
    "Facade photos and generic material textures were placed under `assets/textures/`.",
    "The final runtime scene is loaded from one main campus file plus a few extra landmark models.",
]:
    add_bullet(item)
add_paragraph(
    "This workflow saved a lot of time. It let us focus on rendering and interaction in C++ "
    "instead of hand-building every campus object inside the code."
)

add_heading("6. Rendering Pipeline")
add_paragraph(
    "The renderer is a forward pipeline. Each frame, the program updates the camera, draws the "
    "skybox, uploads the current light and camera uniforms, and then renders each visible mesh."
)
add_paragraph(
    "We kept the draw logic simple on purpose. The scene is static, so our main work was to "
    "load the geometry once, keep per-mesh metadata, and then render only what is needed."
)
add_code(
    "display();\n"
    "drawSkybox(view, projection, lightPreset);\n"
    "setUniform(shaderProgram, \"view\", view);\n"
    "setUniform(shaderProgram, \"projection\", projection);\n"
    "drawModel(campusModel);\n"
    "drawModel(odeonModel);\n"
    "drawModel(clockTowerModel);\n"
    "drawHud(view, projection);"
)
add_paragraph(
    "Frustum culling happens before each mesh draw. If a mesh sphere is outside the current "
    "camera frustum, it is skipped. This keeps the final build lighter when the user is close "
    "to only one part of the campus."
)
add_screenshot_placeholder("Figure 2 — Night mode view with the procedural skybox and darker lighting.")

add_heading("7. Camera, Navigation, and Interaction")
add_paragraph(
    "Navigation was one of the most important parts of the user experience. We wanted the demo "
    "to feel easy to control during both a normal walk and a live presentation."
)
control_headers = ["Input", "Action"]
control_rows = [
    ["W / A / S / D", "Move in first-person mode."],
    ["Mouse", "Look around."],
    ["Shift / Alt", "Sprint and turbo movement."],
    ["F", "Toggle fly mode."],
    ["Space / C", "Move up and down in fly mode."],
    ["N", "Toggle day and night lighting."],
    ["M", "Toggle minimap HUD."],
    ["1 / 2 / 3", "Jump to important landmark views."],
]
t = doc.add_table(rows=1 + len(control_rows), cols=2)
t.style = "Table Grid"
style_table(t, control_headers)
for i, row in enumerate(control_rows):
    fill_table_row(t.rows[i + 1], row)

add_paragraph(
    "We also added landmark presets for demo use. This made it easy to move quickly between key "
    "areas without manually flying across the full map every time."
)
add_screenshot_placeholder("Figure 3 — Walk mode near the main buildings, with terrain following and collision.")
add_screenshot_placeholder("Figure 4 — Fly mode used for a high-level campus overview.")

add_heading("8. Terrain Following and Collision")
add_paragraph(
    "Walking on a campus model is not only about moving in X and Z. The camera also needs a "
    "correct Y value, and it should not pass through buildings."
)
add_paragraph(
    "To solve this, we collect walkable triangles from meshes whose names match terrain, road, "
    "or path categories. During movement, the program samples the triangle under the user and "
    "moves the camera to that height. This keeps the player on the ground."
)
add_paragraph(
    "For collision, we build simple obstacle boxes from building bounds. This is cheaper than "
    "triangle-level collision and is enough for our campus tour use case. We also use the same "
    "walk data to find a safe spawn point when the program starts."
)

add_heading("9. Lighting, Materials, and Texture Handling")
add_paragraph(
    "Lighting is handled in GLSL. The fragment shader combines ambient, diffuse, and "
    "Blinn-Phong specular terms. Textures are sampled in linear space, and the result is gamma "
    "corrected again before it is written to the screen."
)
add_paragraph(
    "Not every part of the campus has a custom texture. Because of that, we support two "
    "material paths. Some meshes use diffuse textures loaded from files or embedded assets. "
    "Other meshes use base colors and procedural shading rules inside the fragment shader."
)
for item in [
    "Facade and roof colors are assigned from a natural palette when no texture is available.",
    "Grass, stone, water, fence, foliage, and trunk materials get small shader patterns.",
    "The Rectorate facade has extra procedural detail for windows and trim.",
    "Texture loading uses mipmaps and anisotropic filtering for better distant surfaces.",
]:
    add_bullet(item)

add_heading("10. Extra Visual and Demo Features")
add_paragraph(
    "After the main pipeline was stable, we spent time on features that make the final demo "
    "clearer and more pleasant to present."
)
for item in [
    "Procedural cubemap skybox instead of a flat background.",
    "Day and night presets with different light direction, light color, and ambient tint.",
    "HUD pass for labels, mode text, and a corner minimap.",
    "Landmark markers on both the screen space overlay and the minimap.",
    "Extra detail meshes such as the campus fence, fountain, forest area, and special landmark models.",
]:
    add_bullet(item)
add_screenshot_placeholder("Figure 5 — HUD pass with the minimap and status text.")
add_screenshot_placeholder("Figure 6 — A landmark-focused view used during the final demo.")

add_heading("11. Challenges and Design Decisions")

add_heading("11.1 Why we load meshes at runtime", level=2)
add_paragraph(
    "At first, it would have been possible to build simple campus blocks directly in C++. We "
    "did not choose that path. Loading the scene from exported meshes let us keep the campus "
    "shape closer to the real place and kept `main.cpp` focused on rendering and interaction."
)

add_heading("11.2 Why we flatten transforms during loading", level=2)
add_paragraph(
    "The campus scene is static. Because of that, baking node transforms into vertex positions "
    "during loading is simpler than carrying a deep transform hierarchy every frame."
)

add_heading("11.3 Why we use semantic mesh names", level=2)
add_paragraph(
    "Mesh names became an important part of the pipeline. We use them to decide which surfaces "
    "are walkable, which meshes are buildings, and which objects should receive special "
    "materials. This small naming rule made both movement and shading easier to manage."
)

add_heading("11.4 Why we combine textures with procedural materials", level=2)
add_paragraph(
    "Some parts of the scene have real textures, but not every mesh needed a custom image. A "
    "hybrid approach gave us a better final look. It also kept the asset workload reasonable "
    "for a student project with a full campus scene."
)

add_heading("12. Conclusion")
add_paragraph(
    "The final project delivers the experience we planned at the start of the semester. A user "
    "can open the program, move through a recognisable 3D Koç University campus, inspect it "
    "from the air, and switch between visual modes during a live demo."
)
add_paragraph(
    "Just as important, the project uses the main graphics ideas from the course in a connected "
    "way. We use mesh loading, transformations, camera math, shader programming, lighting, "
    "texture mapping, collision handling, and GPU drawing in one consistent pipeline. The final "
    "result is both a working campus tour and a complete computer graphics project."
)

doc.save(OUTPUT_PATH)
print(f"Saved: {OUTPUT_PATH}")
